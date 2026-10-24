"""
generate_mediaplan.py — media-planner skill 统一入口

用法（库导入）:
    from generate_mediaplan import generate_mediaplan_excel
    generate_mediaplan_excel(data, output_path)

用法（脚本）:
    python3 generate_mediaplan.py --data data.json --out output.xlsx

依赖:
    - openpyxl
    - PyYAML
    - skill 包自带的 excel_engine/ 和 mappings/

路径解析:
    所有内部路径（模板、mapping）都以本文件所在目录为基准解析，
    不依赖外部 cwd 或绝对路径，保证 skill 打包后随处可用。
"""
import os
import sys
import json
from datetime import datetime

# 确保本目录优先于 sys.path（允许 `python3 generate_mediaplan.py` 直接运行）
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import openpyxl  # noqa: E402

from excel_engine.table_writer import TableWriter
from excel_engine.mediaplan_axis import adjust_date_axis
from excel_engine.mediaplan_rows import adjust_data_rows
from excel_engine.mediaplan_phases import apply_phase_merges
from excel_engine.mediaplan_writer import write_ad_sets
from excel_engine.style_helper import safe_write


# ===== 路径常量（相对本文件定位，随 skill 迁移）=====
SKILL_ROOT = os.path.dirname(_HERE)  # skill 根目录
TEMPLATE_PATH = os.path.join(SKILL_ROOT, "references", "MediaPlan_Template_Standard.xlsx")
MAPPINGS_DIR = os.path.join(_HERE, "mappings")


# ===== Sheet 映射清单 =====
SHEET_MAPPINGS = [
    ("cover",          "cover.yaml"),
    ("brief",          "brief.yaml"),
    ("summary",        "summary.yaml"),
    ("strategy",       "strategy.yaml"),
    # mediaplan 是特殊的（schedule 类型，不走 TableWriter）
    ("todo",           "todo.yaml"),
    ("audience",       "audience.yaml"),
    ("landing_utm",    "landing_utm.yaml"),
    ("platform_copy",  "platform_copy.yaml"),
]


def generate_mediaplan_excel(data, output_path, template_path=None, audit_dir=None, include_risks_sheet=True):
    """
    从结构化数据生成完整多 Sheet MediaPlan Excel

    Args:
        data: dict，包含以下 keys（详见 scripts/DATA_SCHEMA.md）:
            - cover: dict            静态 KV
            - brief: dict            静态 KV
            - summary: dict          静态 KV
            - strategy: {"rows": [...]}  3 阶段行列表
            - mediaplan: {
                "campaign_start": str(YYYY-MM-DD),
                "total_days": int,
                "phases": [{"name", "start", "end"}, ...],
                "ad_sets": [{...}, ...]  # A-H 字段 + daily_budget_by_phase
              }
            - todo: {"rows": [...]}
            - audience: {"rows": [...]}
            - landing_utm: {"rows": [...]}
            - platform_copy: {"items": [...]}

        output_path: str  输出 Excel 路径
        template_path: str, optional  覆盖默认模板路径

    Returns:
        output_path
    """
    tpl = template_path or TEMPLATE_PATH
    if not os.path.exists(tpl):
        raise FileNotFoundError(f"Template not found: {tpl}")

    # ── 预算守恒预检（tree v2 结构由渲染函数自管，跳过旧预检）─────────────
    if "mediaplan" in data and "cover" in data and "tree" not in data["mediaplan"]:
        mp = data["mediaplan"]
        cover_budget = data["cover"].get("total_budget")
        if cover_budget is not None:
            from datetime import datetime as _dt, timedelta as _td
            try:
                phases = mp.get("phases", [])
                total_days = mp.get("total_days", 0)
                campaign_start = _dt.strptime(mp["campaign_start"], "%Y-%m-%d")
                # 构建 day_offset → phase_name 映射
                day_phase = {}
                for ph in phases:
                    ps = _dt.strptime(ph["start"], "%Y-%m-%d")
                    pe = _dt.strptime(ph["end"], "%Y-%m-%d")
                    d = ps
                    while d <= pe:
                        offset = (d - campaign_start).days
                        if 0 <= offset < total_days:
                            day_phase[offset] = ph["name"]
                        d += _td(days=1)
                computed_total = 0.0
                reserve_total = 0.0
                for ad_set in mp.get("ad_sets", []):
                    # 预留待确认行不排日预算，金额以 reserved_amount 参与守恒（v1.6 矩阵预留伪平台行）
                    if ad_set.get("budget_field") == "reserved_pending":
                        reserve_total += float(ad_set.get("reserved_amount") or 0)
                        continue
                    bbp = ad_set.get("daily_budget_by_phase", {})
                    for offset in range(total_days):
                        pname = day_phase.get(offset)
                        daily = bbp.get(pname, 0) if pname else 0
                        computed_total += daily
                # 日预算按货币分级取整（矩阵格/天数四舍五入到分）允许尾差，
                # 容差与 strategy_engines 排期互溯门禁一致：max(1, total*1e-4)
                grand_total = computed_total + reserve_total
                tolerance = max(1.0, abs(float(cover_budget)) * 1e-4)
                if abs(grand_total - float(cover_budget)) > tolerance:
                    raise ValueError(
                        f"[预算守恒校验失败] 日预算合计 = {computed_total:.2f}，预留 = {reserve_total:.2f}，"
                        f"合计 {grand_total:.2f}，cover.total_budget = {cover_budget}（容差 {tolerance:.2f}）。"
                        f"请用 strategy_engines.derive_daily_schedule 由预算矩阵反算日预算，"
                        f"确保 daily_budget_by_phase 中每个值是日预算（非阶段总额）。"
                    )
                print(f"[preflight] 预算守恒 ✅  scheduled={computed_total:.2f} + reserve={reserve_total:.2f}"
                      f" == total_budget={cover_budget}")
            except ValueError:
                raise
            except Exception as _e:
                print(f"[warn] 预算守恒预检跳过: {_e}")

    wb = openpyxl.load_workbook(tpl)
    writer = TableWriter.from_workbook(wb)

    # 1-4, 6-9: 通用 TableWriter 处理
    for data_key, mapping_file in SHEET_MAPPINGS:
        if data_key not in data:
            print(f"[skip] {data_key}: 未提供数据")
            continue
        mapping_path = os.path.join(MAPPINGS_DIR, mapping_file)
        try:
            writer.fill_sheet(mapping_path, data[data_key])
        except Exception as e:
            print(f"[warn] {data_key} 写入失败: {e}")

    # 5: MediaPlan Sheet 特殊处理（schedule 类型）
    if "mediaplan" in data:
        mp = data["mediaplan"]
        if "tree" in mp:
            # v2 树状结构：阶段纵向分块，重建整张 ④ sheet
            from excel_engine.mediaplan_tree import render_tree_mediaplan
            render_tree_mediaplan(wb["④ MediaPlan 投放排期"], mp)
        else:
            _fill_mediaplan_sheet(wb, mp)

    # 写入 Cover 交付状态（来自 _meta）
    _write_cover_meta(wb, data.get("_meta", {}))

    # 生成"⑩ 风险与待确认" Sheet（含 KPI 参考区间 + risks）
    # 对客 9-sheet 工作簿传 include_risks_sheet=False；风险/证据表另出内部簿。
    if include_risks_sheet:
        _fill_risks_sheet(wb, data)

    try:
        from openpyxl.workbook.properties import CalcProperties
        if wb.calculation is None:
            wb.calculation = CalcProperties(fullCalcOnLoad=True)
        else:
            wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(output_path)

    # 生成内部审计包（audit.json + audit.md）
    if audit_dir:
        _write_audit_files(data, output_path, audit_dir)

    return output_path


def _fill_mediaplan_sheet(wb, mp_data):
    """④ MediaPlan 投放排期 Sheet 填充（schedule 类型）
    v2.0：支持多平台，按 platform 分组，每组前加标签行。
    """
    ws = wb['④ MediaPlan 投放排期']

    start_str = mp_data["campaign_start"]
    total_days = mp_data["total_days"]
    phases = mp_data["phases"]
    ad_sets = mp_data["ad_sets"]

    # 写 campaign_start 到 B2
    safe_write(ws, 2, 2, datetime.strptime(start_str, "%Y-%m-%d"))

    # v2.0：need_rows = Ad Set 行数 + 平台标签行数
    from excel_engine.mediaplan_writer import calc_total_rows
    total_rows = calc_total_rows(ad_sets)

    # 4 个子模块依次调用
    last_date_col, _ = adjust_date_axis(ws, total_days)
    last_data_row, grand_total_row = adjust_data_rows(ws, total_rows, last_date_col)
    apply_phase_merges(ws, start_str, phases, last_date_col)
    write_ad_sets(ws, ad_sets, start_str, phases, total_days)



def _write_cover_meta(wb, meta: dict):
    """将 _meta 中的交付状态信息写入封面 Sheet。"""
    cover_sheet_names = [s for s in wb.sheetnames if 'Cover' in s or '封面' in s]
    if not cover_sheet_names:
        return
    ws = wb[cover_sheet_names[0]]
    from excel_engine.style_helper import safe_write
    # 写入 B16:C19 区块（对应 cover.yaml v5.2 新增字段）
    status = meta.get("deliverable_status", "draft")
    safe_write(ws, 16, 3, status)  # C16
    safe_write(ws, 17, 3, meta.get("revision", 1))         # C17
    safe_write(ws, 18, 3, meta.get("confirmer", ""))        # C18
    safe_write(ws, 19, 3, meta.get("updated_at", ""))       # C19
    # 标签列
    safe_write(ws, 16, 2, "交付状态 Deliverable Status")
    safe_write(ws, 17, 2, "版本号 Revision")
    safe_write(ws, 18, 2, "确认人 Confirmer")
    safe_write(ws, 19, 2, "最后更新 Updated At")
    # draft 状态在 status 格加视觉提示
    cell = ws.cell(row=16, column=3)
    if status == "draft":
        cell.value = "⚠️ 草案 DRAFT"
    elif status == "final":
        cell.value = "✅ 正式版 FINAL"


def _fill_risks_sheet(wb, data: dict):
    """
    生成"⑩ 风险与待确认"Sheet。
    包含：KPI 参考区间 + KPI Recommendation + _risks 风险项。
    不对客展示置信评分细节，只展示风险等级、项目、问题/假设、建议动作、责任人、状态。
    """
    # 若 Sheet 不存在则新建
    sheet_name = "⑩ 风险与待确认"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # 清空现有内容
        for row in ws.iter_rows():
            for cell in row:
                cell.value = None
    else:
        ws = wb.create_sheet(sheet_name)

    from excel_engine.style_helper import safe_write

    # ── 表头 ──
    headers = ["等级", "类型", "项目 / 指标", "问题 / 假设", "建议动作", "责任人", "状态"]
    for col_idx, h in enumerate(headers, start=2):  # B 列起
        safe_write(ws, 1, col_idx, h)

    row = 2

    # ── KPI 参考区间（来自 _kpi）──
    for kpi_item in data.get("_kpi", {}).get("items", []):
        metric = kpi_item.get("metric", "")
        forecast = kpi_item.get("kpi_forecast", {})
        rec = kpi_item.get("kpi_recommendation")

        if not forecast:
            continue

        kpi_status = forecast.get("kpi_forecast_status")
        admissibility = forecast.get("admissibility", "")
        lower = forecast.get("lower_bound")
        upper = forecast.get("upper_bound")

        if kpi_status == "insufficient_evidence":
            problem = "暂无同场景可靠基准，不提供确定性预测"
            action = "建议投放后用实际数据验证"
            level = "WARN"
        elif admissibility == "directional_only":
            ref_range = f"[{lower}, {upper}]" if lower is not None else "—"
            evidence_refs = ", ".join(forecast.get("evidence_refs", []))
            problem = f"参考假设区间 {ref_range}（仅方向性，不得作为正式承诺）"
            action = f"来源: {evidence_refs or '—'}；建议投放后验证"
            level = "WARN"
        elif admissibility == "forecast_allowed":
            ref_range = f"[{lower}, {upper}]" if lower is not None else "—"
            evidence_refs = ", ".join(forecast.get("evidence_refs", []))
            problem = f"预测区间 {ref_range}"
            action = f"来源: {evidence_refs or '—'}"
            level = "INFO"
        else:
            continue

        safe_write(ws, row, 2, level)
        safe_write(ws, row, 3, "KPI 参考区间")
        safe_write(ws, row, 4, metric)
        safe_write(ws, row, 5, problem)
        safe_write(ws, row, 6, action)
        safe_write(ws, row, 7, "")
        safe_write(ws, row, 8, "待确认")
        row += 1

        # KPI Recommendation（独立行）
        if rec:
            suggested = rec.get("suggested_value", "")
            reason = rec.get("reason", "")
            safe_write(ws, row, 2, "WARN")
            safe_write(ws, row, 3, "KPI 优化建议")
            safe_write(ws, row, 4, metric)
            safe_write(ws, row, 5, f"建议值: {suggested}｜{reason}")
            safe_write(ws, row, 6, "人工确认是否接受建议（不影响 KPI Target）")
            safe_write(ws, row, 7, "")
            safe_write(ws, row, 8, "待确认")
            row += 1

    # ── _risks 风险项 ──
    for risk in data.get("_risks", {}).get("items", []):
        safe_write(ws, row, 2, risk.get("severity", "WARN"))
        safe_write(ws, row, 3, "风险项")
        safe_write(ws, row, 4, risk.get("field", ""))
        safe_write(ws, row, 5, risk.get("description", ""))
        safe_write(ws, row, 6, risk.get("action", ""))
        safe_write(ws, row, 7, risk.get("confirmed_by") or "—")
        safe_write(ws, row, 8, "待确认" if not risk.get("confirmed_by") else "已确认")
        row += 1

    if row == 2:
        safe_write(ws, 2, 2, "（无风险项）")


def _write_audit_files(data: dict, excel_path: str, audit_dir: str):
    """
    生成内部审计包：
      {project_id}_r{revision}_audit.json  — 字段状态/证据/置信/校验/责任节点
      {project_id}_r{revision}_audit.md    — 人工可读摘要

    对客只发 Excel，审计包仅内部使用。
    """
    import os, json
    from datetime import datetime

    meta = data.get("_meta", {})
    project_id = meta.get("project_id", "unknown")
    revision = meta.get("revision", 1)
    base_name = f"{project_id}_r{revision}"

    os.makedirs(audit_dir, exist_ok=True)

    # audit.json：提取内部底稿层
    audit_data = {
        "_meta":      meta,
        "_field_meta": data.get("_field_meta", {}),
        "_kpi":       data.get("_kpi", {}),
        "_evidence":  data.get("_evidence", {}),
        "_risks":     data.get("_risks", {}),
        "_generated_at": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "_excel_path": excel_path,
    }
    audit_json_path = os.path.join(audit_dir, f"{base_name}_audit.json")
    with open(audit_json_path, "w", encoding="utf-8") as f:
        json.dump(audit_data, f, ensure_ascii=False, indent=2, default=str)

    # audit.md：人工可读摘要
    lines = [
        f"# 审计摘要 — {project_id} r{revision}",
        "",
        f"- **交付状态**: {meta.get('deliverable_status', 'draft')}",
        f"- **确认人**: {meta.get('confirmer', '—')}",
        f"- **更新时间**: {meta.get('updated_at', '—')}",
        f"- **Excel**: `{os.path.basename(excel_path)}`",
        "",
        "## KPI 参考区间",
        "",
    ]
    for item in data.get("_kpi", {}).get("items", []):
        metric = item.get("metric", "")
        forecast = item.get("kpi_forecast", {})
        if forecast:
            kpi_status = forecast.get("kpi_forecast_status", "")
            if kpi_status == "insufficient_evidence":
                lines.append(f"- **{metric}**: 暂无可靠基准")
            else:
                lb = forecast.get("lower_bound", "—")
                ub = forecast.get("upper_bound", "—")
                adm = forecast.get("admissibility", "—")
                lines.append(f"- **{metric}**: [{lb}, {ub}]  admissibility={adm}")
    lines += ["", "## 风险项", ""]
    for risk in data.get("_risks", {}).get("items", []):
        lines.append(
            f"- [{risk.get('severity','WARN')}] {risk.get('field','')} — "
            f"{risk.get('description','')} | 动作: {risk.get('action','')}"
        )
    if not data.get("_risks", {}).get("items"):
        lines.append("（无风险项）")

    lines += ["", "## 证据记录摘要", ""]
    for eid, rec in data.get("_evidence", {}).get("records", {}).items():
        cs = rec.get("confidence_score", {})
        lines.append(
            f"- `{eid}`: {rec.get('source_ref','—')} | "
            f"tier={cs.get('tier','?')} admissibility={cs.get('admissibility','?')}"
        )
    if not data.get("_evidence", {}).get("records"):
        lines.append("（无证据记录）")

    audit_md_path = os.path.join(audit_dir, f"{base_name}_audit.md")
    with open(audit_md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    return audit_json_path, audit_md_path

# ===== CLI =====

def _main():
    import argparse
    parser = argparse.ArgumentParser(description="生成广告 MediaPlan Excel")
    parser.add_argument("--data", required=True, help="数据 JSON 文件路径")
    parser.add_argument("--out", required=True, help="输出 Excel 路径")
    parser.add_argument("--template", help="可选：覆盖模板路径")
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    out = generate_mediaplan_excel(data, args.out, template_path=args.template)
    print(f"✅ Generated: {out}")


if __name__ == "__main__":
    _main()

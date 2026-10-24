"""
TableWriter - 核心填表引擎 (P4 新增 grouped_list_row)

支持的 table_type:
  - static_kv
  - list_row
  - list_col
  - grouped_list_row (P4)
"""
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .mapping_loader import load_mapping, validate_template_against_mapping
from .style_helper import safe_write, copy_row_style, copy_col_style, copy_cell_style


class TableWriter:
    def __init__(self, template_path, strict=False):
        self.template_path = template_path
        if template_path is not None:
            self.wb = openpyxl.load_workbook(template_path)
        else:
            self.wb = None  # set later via attach_wb
        self.strict = strict
        self.warnings = []

    @classmethod
    def from_workbook(cls, wb, strict=False):
        """复用已加载的 workbook（P6 组装时用）"""
        inst = cls(None, strict=strict)
        inst.wb = wb
        return inst

    def fill_sheet(self, mapping_path, data):
        mapping = load_mapping(mapping_path)
        sheet_name = mapping["sheet_name"]
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not in workbook")
        ws = self.wb[sheet_name]

        table_type = mapping["table_type"]

        # 模板校验
        if table_type == "static_kv":
            ok, warnings_ = validate_template_against_mapping(ws, mapping)
            if not ok:
                if self.strict:
                    raise RuntimeError("\n".join(warnings_))
                self.warnings.extend(warnings_)
        elif table_type in ("list_row", "list_col", "grouped_list_row"):
            hv = mapping.get("header_validate", {})
            for cell_ref, expected in hv.items():
                actual = ws[cell_ref].value
                if actual != expected:
                    msg = f"Sheet '{sheet_name}' header {cell_ref} expected '{expected}', got '{actual}'"
                    if self.strict:
                        raise RuntimeError(msg)
                    self.warnings.append(msg)

        # 分发
        if table_type == "static_kv":
            self._fill_static_kv(ws, mapping, data)
        elif table_type == "list_row":
            self._fill_list_row(ws, mapping, data)
        elif table_type == "list_col":
            self._fill_list_col(ws, mapping, data)
        elif table_type == "grouped_list_row":
            self._fill_grouped_list_row(ws, mapping, data)
        else:
            raise NotImplementedError(f"table_type '{table_type}' not supported yet")
        return self

    # ========================================================
    # static_kv
    # ========================================================
    def _fill_static_kv(self, ws, mapping, data):
        for field_key, field_def in mapping.get("fields", {}).items():
            value_cell = field_def.get("value_cell")
            if field_key in data and data[field_key] is not None:
                val = data[field_key]
            elif field_def.get("fallback_field") and data.get(field_def["fallback_field"]) is not None:
                val = data[field_def["fallback_field"]]
            elif "default" in field_def:
                val = field_def["default"]
            elif field_def.get("auto") == "today":
                val = datetime.now().strftime("%Y-%m-%d")
            else:
                continue
            cell = ws[value_cell]
            safe_write(ws, cell.row, cell.column, val)

    # ========================================================
    # list_row
    # ========================================================
    def _fill_list_row(self, ws, mapping, data):
        rows = data.get("rows", [])
        if not isinstance(rows, list):
            raise ValueError("list_row data must contain 'rows' as a list")

        data_start = mapping["data_start_row"]
        style_src = mapping.get("style_source_row", data_start)
        columns = mapping["columns"]

        # 预处理：解除合并
        if mapping.get("pre_unmerge"):
            for mr in list(ws.merged_cells.ranges):
                if mr.min_row >= data_start:
                    ws.unmerge_cells(str(mr))

        # 模板预留行数
        if "template_reserved_rows" in mapping:
            reserved_rows = mapping["template_reserved_rows"]
        else:
            reserved_rows = 0
            scan_cols = [column_index_from_string(c["col"]) for c in columns.values()][:3]
            was_empty_colored = False
            for r in range(data_start, ws.max_row + 2):
                has_content = any(ws.cell(row=r, column=c).value not in (None, "") for c in scan_cols)
                # Check if row has a non-default fill (pre-colored placeholder row)
                has_fill = False
                for c in scan_cols:
                    cell = ws.cell(row=r, column=c)
                    if cell.fill and cell.fill.patternType:
                        rgb = getattr(cell.fill.start_color, "rgb", None)
                        if rgb and isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF", None):
                            has_fill = True
                            break
                if not has_content and not has_fill:
                    break  # truly blank/unstyled row — end of template area
                # If we were counting empty colored placeholder rows and now hit a
                # content row, that content row is likely the Total/summary row — stop.
                if has_content and was_empty_colored:
                    break
                reserved_rows += 1
                was_empty_colored = was_empty_colored or (not has_content and has_fill)
                if reserved_rows > 200:
                    break
            reserved_rows = max(reserved_rows, 1)

        need_rows = len(rows)

        # 行数调整
        if need_rows > reserved_rows:
            insert_at = data_start + reserved_rows
            extra = need_rows - reserved_rows
            ws.insert_rows(insert_at, extra)
            for i in range(extra):
                tgt_row = insert_at + i
                max_col = max(column_index_from_string(c["col"]) for c in columns.values())
                copy_row_style(ws, style_src, tgt_row, 1, max_col)
        elif need_rows < reserved_rows and mapping.get("cleanup_empty_rows", {}).get("enabled"):
            min_keep = mapping["cleanup_empty_rows"].get("min_keep", 0)
            actual_keep = max(need_rows, min_keep)
            to_delete = reserved_rows - actual_keep
            if to_delete > 0:
                delete_at = data_start + actual_keep
                ws.delete_rows(delete_at, to_delete)

        # 写入
        for i, row_data in enumerate(rows):
            r = data_start + i
            for col_key, col_def in columns.items():
                col_letter = col_def["col"]
                field = col_def["field"]
                fallback = col_def.get("fallback_field")
                col_idx = column_index_from_string(col_letter)
                val = row_data.get(field)
                if val is None and fallback:
                    val = row_data.get(fallback)
                if val is not None:
                    safe_write(ws, r, col_idx, val)
                    if col_def.get("number_format"):
                        ws.cell(row=r, column=col_idx).number_format = col_def["number_format"]

        # ── 后处理：插入的额外阶段行分配区分色 ──
        if need_rows > reserved_rows:
            from openpyxl.styles import PatternFill
            EXTRA_PHASE_FILLS = ["FFDDEBF7", "FFFFE699", "FFE4DFEC", "FFD9EAD3"]
            max_col = max(column_index_from_string(c["col"]) for c in columns.values())
            for i in range(reserved_rows, need_rows):
                tgt_row = data_start + i
                color_rgb = EXTRA_PHASE_FILLS[(i - reserved_rows) % len(EXTRA_PHASE_FILLS)]
                fill = PatternFill("solid", fgColor=color_rgb)
                for c in range(1, max_col + 1):
                    ws.cell(row=tgt_row, column=c).fill = fill

        # ── 后处理：百分比列总计写成数值（不写公式，外部查看器不空白）──
        pct_col_idx = None
        for col_def in columns.values():
            if col_def.get("number_format") == "0.0%":
                pct_col_idx = column_index_from_string(col_def["col"])
                break
        if pct_col_idx is not None:
            total_sum = 0.0
            for i in range(need_rows):
                v = ws.cell(row=data_start + i, column=pct_col_idx).value
                if isinstance(v, (int, float)):
                    total_sum += v
            search_row = data_start + need_rows
            for offset in range(0, 5):
                r = search_row + offset
                row_text = " ".join(
                    str(ws.cell(row=r, column=c).value or "")
                    for c in range(1, pct_col_idx + 1)
                )
                if "合计" in row_text or "Total" in row_text or "total" in row_text:
                    cell = ws.cell(row=r, column=pct_col_idx)
                    cell.value = round(total_sum, 6)
                    cell.number_format = "0.0%"
                    break

        # ── 自动行高：按文字长度和列宽估算，避免内容被截断 ──
        self._autofit_row_heights(ws, data_start, data_start + need_rows - 1, columns)

    def _autofit_row_heights(self, ws, row_start, row_end, columns):
        import math
        LINE_HEIGHT = 15.0
        MIN_HEIGHT = 22.0
        for r in range(row_start, row_end + 1):
            max_lines = 1
            for col_def in columns.values():
                col_letter = col_def["col"]
                col_idx = column_index_from_string(col_letter)
                cell = ws.cell(row=r, column=col_idx)
                val = cell.value
                if not val:
                    continue
                text = str(val)
                col_width = ws.column_dimensions[col_letter].width or 10
                eff_chars_per_line = max(5, int(col_width * 0.9))
                for seg in text.split("\n"):
                    seg_len = sum(2 if ord(ch) > 127 else 1 for ch in seg)
                    lines_needed = math.ceil(seg_len / (eff_chars_per_line * 1.2))
                    max_lines = max(max_lines, lines_needed)
            needed_height = max(MIN_HEIGHT, max_lines * LINE_HEIGHT + 6)
            ws.row_dimensions[r].height = needed_height

    # ========================================================
    # grouped_list_row - 分组行列表（带合并重建）
    # ========================================================
    def _fill_grouped_list_row(self, ws, mapping, data):
        # 复用 list_row 的填充逻辑
        self._fill_list_row(ws, mapping, data)

        # 重建合并
        rows = data.get("rows", [])
        if not rows:
            return
        data_start = mapping["data_start_row"]
        need_rows = len(rows)
        last_row = data_start + need_rows - 1
        columns = mapping["columns"]

        for rule in mapping.get("merge_rules", []):
            col_letter = rule["col"]
            col_idx = column_index_from_string(col_letter)
            strategy = rule["strategy"]

            if strategy == "consecutive_same":
                self._merge_consecutive_same(ws, col_idx, data_start, last_row)
            elif strategy == "group_prefix":
                self._merge_group_prefix(ws, col_idx, data_start, last_row,
                                         rule.get("prefix_rules", []))

    def _merge_consecutive_same(self, ws, col_idx, row_start, row_end):
        """连续相同值的单元格合并"""
        i = row_start
        while i <= row_end:
            val = ws.cell(row=i, column=col_idx).value
            if val in (None, ""):
                i += 1
                continue
            j = i
            while j + 1 <= row_end and ws.cell(row=j + 1, column=col_idx).value == val:
                j += 1
            if j > i:
                # 合并 i -> j
                rng = f"{get_column_letter(col_idx)}{i}:{get_column_letter(col_idx)}{j}"
                ws.merge_cells(rng)
                # 清空后续单元格值（openpyxl 合并会自动保留左上角）
            i = j + 1

    def _merge_group_prefix(self, ws, col_idx, row_start, row_end, prefix_rules):
        """按前缀分组合并（如 B1/B2/B3/B4 → 合并为 B）
        保留左上角的完整样式（fill/font/border/alignment）"""
        from copy import copy as _copy
        for rule in prefix_rules:
            group_label = rule["group_label"]
            applies_to = set(rule["applies_to"])

            # 找连续属于该分组的行段
            segments = []
            seg_start = None
            for r in range(row_start, row_end + 1):
                val = ws.cell(row=r, column=col_idx).value
                if val in applies_to:
                    if seg_start is None:
                        seg_start = r
                    seg_end = r
                else:
                    if seg_start is not None:
                        segments.append((seg_start, seg_end))
                        seg_start = None
            if seg_start is not None:
                segments.append((seg_start, seg_end))

            for s, e in segments:
                anchor = ws.cell(row=s, column=col_idx)
                # 抓取样式快照
                snapshot = {
                    "font": _copy(anchor.font),
                    "border": _copy(anchor.border),
                    "fill": _copy(anchor.fill),
                    "alignment": _copy(anchor.alignment),
                    "number_format": anchor.number_format,
                    "protection": _copy(anchor.protection),
                }
                # 先写值
                safe_write(ws, s, col_idx, group_label)
                # 再恢复样式
                anchor_after = ws.cell(row=s, column=col_idx)
                anchor_after.font = snapshot["font"]
                anchor_after.border = snapshot["border"]
                anchor_after.fill = snapshot["fill"]
                anchor_after.alignment = snapshot["alignment"]
                anchor_after.number_format = snapshot["number_format"]
                anchor_after.protection = snapshot["protection"]

                if e > s:
                    rng = f"{get_column_letter(col_idx)}{s}:{get_column_letter(col_idx)}{e}"
                    ws.merge_cells(rng)

    # ========================================================
    # list_col
    # ========================================================
    def _fill_list_col(self, ws, mapping, data):
        items = data.get("items", [])
        if not isinstance(items, list):
            raise ValueError("list_col data must contain 'items' as a list")

        data_start_col = column_index_from_string(mapping["data_start_col"])
        style_src_col = column_index_from_string(mapping.get("style_source_col", mapping["data_start_col"]))
        rows_map = mapping["rows"]

        if "template_reserved_cols" in mapping:
            reserved_cols = mapping["template_reserved_cols"]
        else:
            scan_rows = [r["row"] for r in rows_map.values()]
            reserved_cols = 0
            for c in range(data_start_col, ws.max_column + 2):
                has_content = False
                for r in scan_rows:
                    cell = ws.cell(row=r, column=c)
                    if cell.value not in (None, ""):
                        has_content = True
                        break
                    if cell.has_style and cell.fill and cell.fill.start_color:
                        rgb = getattr(cell.fill.start_color, "rgb", None)
                        if rgb and isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF", None):
                            has_content = True
                            break
                if not has_content:
                    break
                reserved_cols += 1
                if reserved_cols > 50:
                    break
            reserved_cols = max(reserved_cols, 1)

        need_cols = len(items)

        if need_cols > reserved_cols:
            insert_at = data_start_col + reserved_cols
            extra = need_cols - reserved_cols
            ws.insert_cols(insert_at, extra)
            for i in range(extra):
                tgt_col = insert_at + i
                max_row = max(rows_map[k]["row"] for k in rows_map)
                copy_col_style(ws, style_src_col, tgt_col, 1, max_row)
        elif need_cols < reserved_cols and mapping.get("cleanup_empty_cols", {}).get("enabled"):
            min_keep = mapping["cleanup_empty_cols"].get("min_keep", 0)
            actual_keep = max(need_cols, min_keep)
            to_delete = reserved_cols - actual_keep
            if to_delete > 0:
                delete_at = data_start_col + actual_keep
                ws.delete_cols(delete_at, to_delete)

        for i, item in enumerate(items):
            c = data_start_col + i
            for row_key, row_def in rows_map.items():
                r = row_def["row"]
                field = row_def["field"]
                fallback = row_def.get("fallback_field")
                if field in item and item[field] is not None:
                    val = item[field]
                elif fallback and item.get(fallback) is not None:
                    val = item[fallback]
                elif "default" in row_def:
                    val = row_def["default"]
                else:
                    continue
                safe_write(ws, r, c, val)

    def save(self, output_path):
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output_path)
        return output_path

"""
样式辅助工具 - PoC 引擎的地基

职责：
1. 单元格级样式复制（font/border/fill/alignment/number_format/protection）
2. 行/列几何属性复制（row_dimensions.height, column_dimensions.width）
3. 合并单元格范围复制
4. 安全写入（MergedCell 自动解合并）
"""
from copy import copy
from openpyxl.cell.cell import Cell
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.utils import get_column_letter, column_index_from_string


# =========================================================
# 1. 单元格样式复制
# =========================================================
def copy_cell_style(src_cell, tgt_cell):
    """把 src_cell 的所有样式属性复制到 tgt_cell"""
    if src_cell.has_style:
        tgt_cell.font = copy(src_cell.font)
        tgt_cell.border = copy(src_cell.border)
        tgt_cell.fill = copy(src_cell.fill)
        tgt_cell.alignment = copy(src_cell.alignment)
        tgt_cell.number_format = copy(src_cell.number_format)
        tgt_cell.protection = copy(src_cell.protection)


def copy_cell(ws, src_ref, tgt_ref, with_value=False):
    """
    按单元格引用复制样式（可选连值）
    src_ref/tgt_ref: 支持 'A1' 字符串 或 (row, col) 元组
    """
    src = ws[src_ref] if isinstance(src_ref, str) else ws.cell(row=src_ref[0], column=src_ref[1])
    tgt = ws[tgt_ref] if isinstance(tgt_ref, str) else ws.cell(row=tgt_ref[0], column=tgt_ref[1])
    copy_cell_style(src, tgt)
    if with_value:
        tgt.value = src.value
    return tgt


# =========================================================
# 2. 行/列几何属性复制
# =========================================================
def copy_column_width(ws, src_col, tgt_col):
    """src_col/tgt_col 支持整数（1-based）或字母（'A'）"""
    src_letter = src_col if isinstance(src_col, str) else get_column_letter(src_col)
    tgt_letter = tgt_col if isinstance(tgt_col, str) else get_column_letter(tgt_col)
    if src_letter in ws.column_dimensions:
        ws.column_dimensions[tgt_letter].width = ws.column_dimensions[src_letter].width


def copy_row_height(ws, src_row, tgt_row):
    """行高复制，行号 1-based"""
    if src_row in ws.row_dimensions:
        ws.row_dimensions[tgt_row].height = ws.row_dimensions[src_row].height


# =========================================================
# 3. 合并单元格范围复制
# =========================================================
def copy_merged_range(ws, src_range, row_offset=0, col_offset=0):
    """
    复制一个合并范围到 (row_offset, col_offset) 偏移处
    src_range: 'A1:C3' 格式
    """
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(src_range)
    new_min_row = min_row + row_offset
    new_max_row = max_row + row_offset
    new_min_col = min_col + col_offset
    new_max_col = max_col + col_offset
    new_range = f"{get_column_letter(new_min_col)}{new_min_row}:{get_column_letter(new_max_col)}{new_max_row}"
    ws.merge_cells(new_range)
    # 复制左上角样式
    copy_cell_style(ws.cell(row=min_row, column=min_col),
                    ws.cell(row=new_min_row, column=new_min_col))


# =========================================================
# 4. 安全写入
# =========================================================
def safe_write(ws, row, col, value):
    """
    写入单元格，如果遇到 MergedCell 则自动解合并再写
    保留解合并前的样式
    """
    try:
        ws.cell(row=row, column=col, value=value)
    except AttributeError:
        # MergedCell 场景：找到所在的合并范围并解除
        for mr in list(ws.merged_cells.ranges):
            if (mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col):
                # 先记录左上角样式
                anchor = ws.cell(row=mr.min_row, column=mr.min_col)
                anchor_style = {
                    "font": copy(anchor.font),
                    "border": copy(anchor.border),
                    "fill": copy(anchor.fill),
                    "alignment": copy(anchor.alignment),
                    "number_format": copy(anchor.number_format),
                    "protection": copy(anchor.protection),
                }
                ws.unmerge_cells(str(mr))
                # 写入值
                ws.cell(row=row, column=col, value=value)
                # 恢复左上角样式（其他单元格保留默认即可）
                anchor = ws.cell(row=mr.min_row, column=mr.min_col)
                for k, v in anchor_style.items():
                    setattr(anchor, k, v)
                return
        # 未匹配到合并范围，原样抛出
        raise


# =========================================================
# 5. 行/列整体复制（样式，不含值）
# =========================================================
def copy_row_style(ws, src_row, tgt_row, col_start=1, col_end=None):
    """把 src_row 的样式复制到 tgt_row，列范围 [col_start, col_end]"""
    if col_end is None:
        col_end = ws.max_column
    for c in range(col_start, col_end + 1):
        src = ws.cell(row=src_row, column=c)
        tgt = ws.cell(row=tgt_row, column=c)
        copy_cell_style(src, tgt)
    copy_row_height(ws, src_row, tgt_row)


def copy_col_style(ws, src_col, tgt_col, row_start=1, row_end=None):
    """把 src_col 的样式复制到 tgt_col，行范围 [row_start, row_end]"""
    if row_end is None:
        row_end = ws.max_row
    src_idx = src_col if isinstance(src_col, int) else column_index_from_string(src_col)
    tgt_idx = tgt_col if isinstance(tgt_col, int) else column_index_from_string(tgt_col)
    for r in range(row_start, row_end + 1):
        src = ws.cell(row=r, column=src_idx)
        tgt = ws.cell(row=r, column=tgt_idx)
        copy_cell_style(src, tgt)
    copy_column_width(ws, src_idx, tgt_idx)

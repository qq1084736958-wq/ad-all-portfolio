# -*- coding: utf-8 -*-
"""
mediaplan_tree.py — 树状 MediaPlan 渲染（v2 新结构）

列布局（1-based）：
  A 阶段 | B Platform | C Nation | D Ad Format | E Goal | F Product |
  G Targeting | H Placement | I Creative | J Line Total | K Budget(日预算) | L.. 日期列

行布局（纵向分块）：
  R1 标题 / R2 Campaign Start / R3 日期 / R4 星期 / R5 列头
  R6+ 按阶段纵向分块：
    阶段小计行（A 列阶段名，J 列阶段小计）
      owner 小计行（B 列 "XXX：合计"）
        明细行（拆到 platform/goal/product/targeting/creative）
"""
from datetime import datetime, timedelta
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FILL_TITLE = PatternFill("solid", fgColor="FF4472C4")
FILL_PHASE = PatternFill("solid", fgColor="FF1F4E79")
FILL_HDR = PatternFill("solid", fgColor="FFD9E1F2")

_OWNER_FILL = {
    "meta": PatternFill("solid", fgColor="FFFCE4D6"),
    "google": PatternFill("solid", fgColor="FFDDEBF7"),
    "youtube": PatternFill("solid", fgColor="FFDDEBF7"),
    "tiktok": PatternFill("solid", fgColor="FFE2EFDA"),
}
_FILL_OWNER_DEFAULT = PatternFill("solid", fgColor="FFFFF2CC")


def _owner_fill(owner):
    key = (owner or "").strip().casefold()
    if key in _OWNER_FILL:
        return _OWNER_FILL[key]
    for name, fill in _OWNER_FILL.items():
        if name in key:
            return fill
    return _FILL_OWNER_DEFAULT

FONT_BASE = Font(name="Microsoft YaHei", size=11)
FONT_BOLD = Font(name="Microsoft YaHei", size=11, bold=True)
FONT_WHITE_BOLD = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFFFF")
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COL_PHASE, COL_PLATFORM, COL_NATION, COL_FORMAT, COL_GOAL = 1, 2, 3, 4, 5
COL_PRODUCT, COL_TARGET, COL_PLACE, COL_CREATIVE = 6, 7, 8, 9
COL_LINETOTAL, COL_BUDGET = 10, 11
COL_DATE_START = 12

PHASE_NAME = {
    "P1": "阶段1 P1 线上预热", "P2": "阶段2 P2 双渠道预售",
    "P3": "阶段3 P3 开售爆发", "P4": "阶段4 P4 首销长尾",
}


def _style(cell, font=FONT_BASE, fill=None, align=LEFT, border=BORDER):
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align
    cell.border = border


def render_tree_mediaplan(ws, mp):
    start = datetime.strptime(mp["campaign_start"], "%Y-%m-%d")
    total_days = mp["total_days"]
    phases = mp["phases"]
    tree = mp.get("tree", [])

    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    ws.delete_rows(1, ws.max_row)

    last_date_col = COL_DATE_START + total_days - 1

    ws.cell(1, 1, "Media Plan 投放排期（树状：阶段→平台→Goal→产品→人群→素材）")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_date_col)
    c = ws.cell(1, 1)
    c.font = FONT_WHITE_BOLD; c.fill = FILL_TITLE; c.alignment = CENTER
    ws.row_dimensions[1].height = 26

    ws.cell(2, 1, "Campaign Start"); ws.cell(2, 1).font = FONT_BOLD
    dc = ws.cell(2, 2, start); dc.number_format = "yyyy-mm-dd"
    dc.font = FONT_BASE
    ws.cell(2, COL_DATE_START - 1, "Campaign Period 投放周期").font = FONT_BOLD
    ws.cell(2, COL_DATE_START, f'{mp["campaign_start"]} ~ {(start+timedelta(days=total_days-1)).date()}').font = FONT_BASE

    for off in range(total_days):
        col = COL_DATE_START + off
        d = start + timedelta(days=off)
        c3 = ws.cell(3, col, d); c3.number_format = "m/d"; _style(c3, FONT_BOLD, FILL_HDR, CENTER)
        c4 = ws.cell(4, col, d.strftime("%a")); _style(c4, FONT_BASE, FILL_HDR, CENTER)
    for lbl_row, lbl in [(3, "日期 Date"), (4, "星期 Weekday")]:
        for col in range(1, COL_DATE_START):
            _style(ws.cell(lbl_row, col, ""), FONT_BASE, FILL_HDR, CENTER)

    headers = ["阶段 Phase", "Platform 平台", "Nation 国家", "Ad Format 广告形式", "Goal 目标",
               "Product 产品", "Targeting 人群", "Placement 版位", "Creative 素材",
               "Line Total 小计", "Budget 日预算"]
    for i, htxt in enumerate(headers, start=1):
        c = ws.cell(5, i, htxt); _style(c, FONT_WHITE_BOLD, FILL_PHASE, CENTER)
    for off in range(total_days):
        _style(ws.cell(5, COL_DATE_START + off, ""), FONT_BASE, FILL_PHASE)
    ws.row_dimensions[5].height = 30

    day_phase = {}
    for off in range(total_days):
        d = (start + timedelta(days=off)).date()
        for p in phases:
            ps = datetime.strptime(p["start"], "%Y-%m-%d").date()
            pe = datetime.strptime(p["end"], "%Y-%m-%d").date()
            if ps <= d <= pe:
                day_phase[off] = p["name"]
                break

    r = 6
    phase_blocks = {}
    for row in tree:
        phase_blocks.setdefault(row["phase"], []).append(row)

    phase_days_map = {}
    for p in phases:
        ps = datetime.strptime(p["start"], "%Y-%m-%d").date()
        pe = datetime.strptime(p["end"], "%Y-%m-%d").date()
        phase_days_map[p["name"]] = (pe - ps).days + 1

    for ph in phases:
        pname = ph["name"]
        rows_in_phase = phase_blocks.get(pname, [])
        if not rows_in_phase:
            continue
        phase_start_row = r
        phase_days = phase_days_map.get(pname, 0)

        pc = ws.cell(r, 1, PHASE_NAME.get(pname, pname))
        pc.font = FONT_WHITE_BOLD; pc.fill = FILL_PHASE; pc.alignment = LEFT; pc.border = BORDER
        for col in range(2, COL_DATE_START):
            cc = ws.cell(r, col, ""); cc.fill = FILL_PHASE; cc.border = BORDER
        jc = ws.cell(r, COL_LINETOTAL)
        jc.font = FONT_WHITE_BOLD; jc.fill = FILL_PHASE; jc.border = BORDER
        bc = ws.cell(r, COL_BUDGET, ""); bc.fill = FILL_PHASE; bc.border = BORDER
        for off in range(total_days):
            cc = ws.cell(r, COL_DATE_START + off, ""); cc.fill = FILL_PHASE; cc.border = BORDER
        phase_summary_row = r
        r += 1

        owner_groups = {}
        for row in rows_in_phase:
            owner_groups.setdefault(row.get("owner", row["channel"]), []).append(row)

        for owner, orows in owner_groups.items():
            owner_fill = _owner_fill(owner)
            oc = ws.cell(r, 2, f"{owner}：合计")
            oc.font = FONT_BOLD; oc.fill = owner_fill; oc.alignment = LEFT; oc.border = BORDER
            for col in [1] + list(range(3, COL_DATE_START)):
                cc = ws.cell(r, col, ""); cc.fill = owner_fill; cc.border = BORDER
            for off in range(total_days):
                cc = ws.cell(r, COL_DATE_START + off, ""); cc.fill = owner_fill; cc.border = BORDER
            owner_summary_row = r
            owner_detail_first = r + 1
            r += 1

            for row in orows:
                vals = {
                    COL_PHASE: "", COL_PLATFORM: row.get("channel", ""),
                    COL_NATION: row.get("nation", "KE"), COL_FORMAT: row.get("ad_format", ""),
                    COL_GOAL: row.get("goal", ""), COL_PRODUCT: row.get("product", ""),
                    COL_TARGET: row.get("targeting", ""), COL_PLACE: row.get("placement", ""),
                    COL_CREATIVE: row.get("creative", ""),
                }
                for col, v in vals.items():
                    _style(ws.cell(r, col, v))
                db = row.get("daily_budget")
                line_total = round(float(db) * phase_days, 2) if isinstance(db, (int, float)) else None
                _style(ws.cell(r, COL_LINETOTAL, line_total), FONT_BOLD, align=CENTER)
                _style(ws.cell(r, COL_BUDGET, db if db not in (None, "") else None), FONT_BASE, align=CENTER)
                for off in range(total_days):
                    cell = ws.cell(r, COL_DATE_START + off)
                    if day_phase.get(off) == pname and db not in (None, ""):
                        cell.value = db
                    _style(cell, FONT_BASE, align=CENTER)
                r += 1
            owner_subtotal = 0.0
            for rr in range(owner_detail_first, r):
                v = ws.cell(rr, COL_LINETOTAL).value
                if isinstance(v, (int, float)):
                    owner_subtotal += v
            ws.cell(owner_summary_row, COL_LINETOTAL, round(owner_subtotal, 2))
            ws.cell(owner_summary_row, COL_LINETOTAL).font = FONT_BOLD
            ws.cell(owner_summary_row, COL_LINETOTAL).fill = owner_fill
            ws.cell(owner_summary_row, COL_LINETOTAL).border = BORDER

        phase_owner_total = 0.0
        rr = phase_start_row + 1
        for owner, orows in owner_groups.items():
            v = ws.cell(rr, COL_LINETOTAL).value
            if isinstance(v, (int, float)):
                phase_owner_total += v
            rr += 1 + len(orows)
        ws.cell(phase_summary_row, COL_LINETOTAL, round(phase_owner_total, 2)).font = FONT_WHITE_BOLD

    widths = {1: 16, 2: 13, 3: 7, 4: 16, 5: 12, 6: 13, 7: 20, 8: 14, 9: 22, 10: 11, 11: 11}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    for off in range(total_days):
        ws.column_dimensions[get_column_letter(COL_DATE_START + off)].width = 9

    return r - 1
